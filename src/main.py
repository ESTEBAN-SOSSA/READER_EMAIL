"""CLI principal: lista solicitudes de cotizacion entrantes y su estado."""
from __future__ import annotations

import logging
import sys
from datetime import datetime, timezone
from pathlib import Path

import pandas as pd
import typer
from rich.console import Console
from rich.logging import RichHandler
from rich.table import Table
from sqlalchemy import delete, desc, select
from sqlalchemy.orm import Session

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except (AttributeError, OSError):
        pass

from .config import load_settings
from .graph_client import GraphClient
from .mail_reader import MailReader, Message
from .storage import (
    ProcessedMessage,
    QuoteRelatedMessage,
    QuoteRequestRow,
    find_by_imid,
    get_skip_ids,
    init_db,
    mark_processed,
    upsert_quote,
    upsert_related,
)
from .summarizer import resumir

app = typer.Typer(
    add_completion=False,
    help="Lector de buzones M365: identifica cotizaciones entrantes y verifica si fueron respondidas.",
)
console = Console()


def _dias_desde(received_at_iso: str) -> int | None:
    """Cuantos dias enteros pasaron desde la fecha recibida hasta hoy."""
    if not received_at_iso:
        return None
    try:
        dt = datetime.fromisoformat(received_at_iso.replace("Z", "+00:00"))
    except ValueError:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    now = datetime.now(timezone.utc)
    delta = now - dt
    return max(delta.days, 0)


def _color_dias(d: int | None) -> str:
    if d is None:
        return "[dim]-[/dim]"
    if d <= 1:
        return f"[green]{d}d[/green]"
    if d <= 3:
        return f"[yellow]{d}d[/yellow]"
    return f"[red bold]{d}d[/red bold]"


def _setup_logging(level: str) -> None:
    logging.basicConfig(
        level=level,
        format="%(message)s",
        datefmt="[%X]",
        handlers=[RichHandler(console=console, rich_tracebacks=True, show_path=False)],
    )
    for noisy in ("httpx", "httpcore", "msal", "urllib3"):
        logging.getLogger(noisy).setLevel(logging.WARNING)


def _guardar(session: Session, msg: Message) -> QuoteRequestRow:
    existente = upsert_quote(session, msg.id)
    if existente:
        # Actualizar estado de respuesta (puede haber cambiado desde la ultima corrida)
        existente.was_replied = msg.was_replied
        existente.replied_at = msg.replied_at
        existente.external_client = msg.external_client
        existente.is_forwarded_internal = msg.is_forwarded_internal
        # Solicitud original (escenario 1): puede detectarse mejor en esta corrida
        existente.original_received_at = msg.original_received_at
        existente.original_snippet = msg.original_snippet
        existente.original_source = msg.original_source
        existente.doc_ref = msg.doc_ref
        existente.group_key = msg.group_key
        existente.is_original = msg.is_original
        existente.original_graph_id = msg.original_graph_id
        existente.internet_message_id = msg.internet_message_id or existente.internet_message_id
        # No degradar 'backfill_via_RE' a 'directo' si ya esta como backfill
        if not existente.chain_source or existente.chain_source == "directo":
            existente.chain_source = msg.chain_source
        session.commit()
        return existente

    resumen = resumir(msg.body_text, [a.name for a in msg.attachments])

    fila = QuoteRequestRow(
        graph_id=msg.id,
        conversation_id=msg.conversation_id,
        mailbox=msg.mailbox,
        subject=msg.subject,
        sender_email=msg.sender_email,
        sender_name=msg.sender_name,
        received_at=msg.received_at,
        snippet=resumen.snippet,
        num_items=resumen.num_items_estimado,
        was_replied=msg.was_replied,
        replied_at=msg.replied_at,
        external_client=msg.external_client,
        is_forwarded_internal=msg.is_forwarded_internal,
        original_received_at=msg.original_received_at,
        original_snippet=msg.original_snippet,
        original_source=msg.original_source,
        doc_ref=msg.doc_ref,
        group_key=msg.group_key,
        is_original=msg.is_original,
        original_graph_id=msg.original_graph_id,
        chain_source=msg.chain_source,
        internet_message_id=msg.internet_message_id or None,
    )
    fila.cantidades = resumen.cantidades_mencionadas
    fila.adjuntos = resumen.adjuntos
    session.add(fila)
    session.commit()
    return fila


@app.command()
def run(
    dry_run: bool = typer.Option(False, "--dry-run", help="No guarda en BD."),
    no_attachments: bool = typer.Option(
        False, "--no-attachments", help="No descargar adjuntos (mas rapido)."
    ),
    reprocess: bool = typer.Option(
        False, "--reprocess", help="Reprocesa todo, ignorando el registro de procesados."
    ),
    no_backfill: bool = typer.Option(
        False, "--no-backfill",
        help="No insertar originales faltantes detectados a partir de una RE.",
    ),
) -> None:
    """Lee buzones, identifica cotizaciones y verifica cuales no han sido respondidas."""
    settings = load_settings()
    _setup_logging(settings.log_level)

    if not settings.mailboxes:
        console.print("[red]No hay MAILBOXES configurados en .env[/red]")
        raise typer.Exit(code=1)

    engine = init_db(settings.database_path)

    with GraphClient(settings.tenant_id, settings.client_id, settings.client_secret) as client:
        reader = MailReader(client, settings)
        with Session(engine) as session:
            skip_ids = set() if reprocess else get_skip_ids(session)
            for mailbox in settings.mailboxes:
                console.rule(f"[bold]Buzon: {mailbox}")
                try:
                    messages, descartados = reader.fetch_relevant_messages(
                        mailbox,
                        download_attachments=not no_attachments,
                        skip_ids=skip_ids,
                    )
                except Exception as exc:
                    console.print(f"[red]Error leyendo {mailbox}: {exc}[/red]")
                    continue

                nuevas = 0
                actualizadas = 0
                backfilled = 0
                re_nuevas = 0
                re_actualizadas = 0
                copias_buzon = 0
                limbo_re: list[str] = []
                backfill_errores: list[tuple[str, str]] = []
                if not dry_run:
                    # Fase B: originales -> quote_requests; respuestas -> sub-tabla.
                    # Primero originales (para que la RE encuentre a su padre).
                    originales = [m for m in messages if m.is_original]
                    respuestas = [m for m in messages if not m.is_original]

                    for msg in originales:
                        existia_por_graph = upsert_quote(session, msg.id) is not None
                        if existia_por_graph:
                            _guardar(session, msg)
                            mark_processed(
                                session, msg.id, mailbox, msg.subject, "identificado"
                            )
                            actualizadas += 1
                            continue
                        # graph_id nuevo: revisar si es duplicado cross-mailbox
                        padre_imid = find_by_imid(session, msg.internet_message_id)
                        if padre_imid is not None:
                            upsert_related(
                                session, msg, padre_imid,
                                relation_type="copia_buzon",
                            )
                            mark_processed(
                                session, msg.id, mailbox, msg.subject,
                                "identificado", "copia_buzon",
                            )
                            copias_buzon += 1
                            continue
                        # Realmente nueva cotizacion
                        _guardar(session, msg)
                        mark_processed(
                            session, msg.id, mailbox, msg.subject, "identificado"
                        )
                        nuevas += 1

                    for msg in respuestas:
                        # Dedup cross-mailbox: misma RE/correo recibido en otro buzon
                        existing_re = session.get(QuoteRelatedMessage, msg.id)
                        if existing_re is None and msg.internet_message_id:
                            padre_imid = find_by_imid(session, msg.internet_message_id)
                            if padre_imid is not None:
                                upsert_related(
                                    session, msg, padre_imid,
                                    relation_type="copia_buzon",
                                )
                                mark_processed(
                                    session, msg.id, mailbox, msg.subject,
                                    "identificado", "copia_buzon",
                                )
                                copias_buzon += 1
                                continue
                        # 1) Localizar al padre (en quote_requests por graph_id)
                        padre = None
                        if msg.original_graph_id:
                            padre = upsert_quote(session, msg.original_graph_id)
                        # 2) Si no esta y backfill activo -> intentar traerlo
                        if padre is None and msg.original_graph_id and not no_backfill:
                            orig_msg = None
                            try:
                                orig_msg = reader.build_original_message(
                                    mailbox,
                                    msg.original_graph_id,
                                    download_attachments=not no_attachments,
                                )
                            except Exception as exc:
                                backfill_errores.append(
                                    (msg.original_graph_id, str(exc)[:80])
                                )
                            if orig_msg:
                                _guardar(session, orig_msg)
                                mark_processed(
                                    session, orig_msg.id, mailbox, orig_msg.subject,
                                    "identificado", "backfill_via_RE",
                                )
                                backfilled += 1
                                nuevas += 1
                                padre = upsert_quote(session, msg.original_graph_id)
                            else:
                                backfill_errores.append(
                                    (msg.original_graph_id, "no_recuperable")
                                )

                        # 3) Sin padre recuperable -> LIMBO, no se inserta en sub-tabla
                        if padre is None:
                            mark_processed(
                                session, msg.id, mailbox, msg.subject,
                                "descartado", "re_sin_padre_limbo",
                            )
                            limbo_re.append(msg.id)
                            continue

                        # 4) Insertar/actualizar en sub-tabla, enlazada al padre
                        es_nueva = upsert_related(session, msg, padre.id)
                        mark_processed(
                            session, msg.id, mailbox, msg.subject,
                            "identificado", "respuesta",
                        )
                        if es_nueva:
                            re_nuevas += 1
                        else:
                            re_actualizadas += 1

                        # 5) Actualizar estado de respuesta del padre con lo que
                        # la RE refleja del hilo (el padre es la fuente de verdad).
                        if msg.was_replied and not padre.was_replied:
                            padre.was_replied = True
                            padre.replied_at = (
                                msg.replied_at or padre.replied_at
                            )
                            padre.external_client = (
                                msg.external_client or padre.external_client
                            )
                            session.commit()

                    # Registrar descartados (filtros) para no reprocesarlos
                    for gid, subj, motivo in descartados:
                        mark_processed(session, gid, mailbox, subj, "descartado", motivo)

                # Reporte
                if backfilled:
                    console.print(
                        f"    [bold cyan]Backfill via RE:[/bold cyan] "
                        f"+{backfilled} originales insertados"
                    )
                if re_nuevas or re_actualizadas:
                    console.print(
                        f"    [bold]Respuestas vinculadas (sub-tabla):[/bold] "
                        f"+{re_nuevas} nuevas, {re_actualizadas} actualizadas"
                    )
                if copias_buzon:
                    console.print(
                        f"    [bold magenta]Copias cross-mailbox:[/bold magenta] "
                        f"{copias_buzon} correo(s) que ya estaban registrados en otro buzon"
                    )
                if backfill_errores:
                    console.print(
                        f"    [yellow]Backfill no recuperable:[/yellow] "
                        f"{len(backfill_errores)} originales"
                    )
                    for gid, err in backfill_errores[:3]:
                        console.print(f"      [dim]- ...{gid[-22:]}: {err}[/dim]")
                if limbo_re:
                    console.print(
                        f"    [yellow bold]REs en LIMBO[/yellow bold] "
                        f"(sin padre recuperable): [bold]{len(limbo_re)}[/bold]"
                    )

                # Resumen consultando la BD (toda la historia acumulada del buzon)
                _mostrar_dashboard_db(
                    session, mailbox, nuevas, actualizadas, len(descartados)
                )

    console.print("\n[bold green]Listo.[/bold green] Usa [cyan]list-pending[/cyan] para ver las pendientes.")


def _mostrar_dashboard_db(
    session: Session, mailbox: str, nuevas: int, actualizadas: int,
    descartados: int = 0,
) -> None:
    """Lee todas las cotizaciones del buzon desde la BD (historico acumulado)."""
    rows = session.execute(
        select(QuoteRequestRow).where(QuoteRequestRow.mailbox == mailbox)
    ).scalars().all()

    total = len(rows)
    if total == 0:
        console.print("  [dim]Sin cotizaciones en historia para este buzon.[/dim]")
        return

    respondidas = sum(1 for r in rows if r.was_replied)
    pendientes_rows = [r for r in rows if not r.was_replied]
    pendientes = len(pendientes_rows)
    pct_resp = (respondidas / total * 100) if total else 0

    dias_pend = [
        d for r in pendientes_rows
        if (d := _dias_desde(r.received_at)) is not None
    ]
    prom_dias = (sum(dias_pend) / len(dias_pend)) if dias_pend else 0
    mas_antigua_dias = max(dias_pend) if dias_pend else 0
    urgentes = sum(1 for d in dias_pend if d > 3)
    criticas = sum(1 for d in dias_pend if d > 7)

    console.print()
    console.print(f"  [bold cyan]Resumen del buzon[/bold cyan] {mailbox}:")
    console.print(
        f"    Esta corrida:  [bold green]+{nuevas}[/bold green] nuevas, "
        f"[dim]{actualizadas} actualizadas, {descartados} descartados[/dim]"
    )
    console.print(f"    [bold]Acumulado en BD[/bold] (toda la historia):")
    console.print(f"      Total identificadas:                 [bold]{total}[/bold]")
    grupos = {r.group_key for r in rows if r.group_key}
    if grupos:
        en_grupos = sum(1 for r in rows if r.group_key)
        console.print(
            f"        [dim]· Agrupadas en {len(grupos)} solicitud(es) consolidada(s) "
            f"({en_grupos} correos)[/dim]"
        )
    # Fase B: respuestas viven en sub-tabla; aqui solo originales
    n_backfill = sum(1 for r in rows if r.chain_source == "backfill_via_RE")
    ids_originales = [r.id for r in rows]
    if ids_originales:
        n_respuestas = len(session.execute(
            select(QuoteRelatedMessage.graph_id).where(
                QuoteRelatedMessage.quote_request_id.in_(ids_originales)
            )
        ).scalars().all())
    else:
        n_respuestas = 0
    if n_respuestas or n_backfill:
        console.print(
            f"        [dim]· {total} cotizaciones (originales) con "
            f"{n_respuestas} correos en sub-tabla "
            f"({n_backfill} originales rescatados por backfill)[/dim]"
        )
    console.print(
        f"      [green]Respondidas al cliente:[/green]   [bold]{respondidas}[/bold]  "
        f"([dim]{pct_resp:.0f}%[/dim])"
    )
    console.print(
        f"      [yellow]PENDIENTES de responder:[/yellow]  [bold]{pendientes}[/bold]  "
        f"([dim]{100 - pct_resp:.0f}%[/dim])"
    )
    if dias_pend:
        console.print(
            f"        [dim]· Espera promedio:[/dim]            "
            f"[bold]{prom_dias:.1f} dias[/bold]"
        )
        console.print(
            f"        [dim]· Mas antigua:[/dim]                "
            f"{_color_dias(mas_antigua_dias)}"
        )
        if urgentes:
            console.print(
                f"        [red]· Mas de 3 dias sin responder:[/red]   "
                f"[bold red]{urgentes}[/bold red]"
            )
        if criticas:
            console.print(
                f"        [bold red]· CRITICAS (>7 dias):[/bold red]            "
                f"[bold red on white] {criticas} [/bold red on white]"
            )


@app.command("list-pending")
def list_pending(limit: int = typer.Option(50, help="Cantidad maxima a mostrar.")) -> None:
    """Lista las cotizaciones que aun NO han sido respondidas."""
    _listar(only_pending=True, limit=limit, titulo="Cotizaciones PENDIENTES de responder")


@app.command("list-all")
def list_all(limit: int = typer.Option(50, help="Cantidad maxima a mostrar.")) -> None:
    """Lista todas las cotizaciones procesadas (respondidas y pendientes)."""
    _listar(only_pending=False, limit=limit, titulo="Todas las cotizaciones procesadas")


def _listar(only_pending: bool, limit: int, titulo: str) -> None:
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    with Session(engine) as session:
        stmt = (
            select(QuoteRequestRow)
            .order_by(desc(QuoteRequestRow.received_at))
            .limit(limit)
        )
        if only_pending:
            stmt = stmt.where(QuoteRequestRow.was_replied.is_(False))
        rows = session.execute(stmt).scalars().all()

    # Ordenar pendientes primero por dias descendentes (las mas viejas arriba)
    rows_sorted = sorted(
        rows,
        key=lambda r: (
            r.was_replied,
            -(_dias_desde(r.received_at) or 0),
        ),
    )

    table = Table(title=titulo, show_lines=True)
    table.add_column("ID", justify="right")
    table.add_column("Dias", justify="right", no_wrap=True)
    table.add_column("Recibido", no_wrap=True)
    table.add_column("Cliente externo", overflow="fold", max_width=24)
    table.add_column("Asunto", overflow="fold", max_width=38)
    table.add_column("Llego de", overflow="fold", max_width=22)
    table.add_column("Estado", no_wrap=True)

    pend = resp = 0
    dias_pendientes: list[int] = []
    for r in rows_sorted:
        if r.was_replied:
            estado = "[green]RESPONDIDA[/green]"
            resp += 1
        else:
            estado = "[yellow]PENDIENTE[/yellow]"
            pend += 1

        dias = _dias_desde(r.received_at)
        if not r.was_replied and dias is not None:
            dias_pendientes.append(dias)

        cliente_txt = r.external_client or "[dim]no identificado[/dim]"
        llego_de = r.sender_email
        if r.is_forwarded_internal:
            llego_de = f"[cyan](reenvio)[/cyan] {llego_de}"

        table.add_row(
            str(r.id),
            _color_dias(dias),
            (r.received_at or "")[:10],
            cliente_txt,
            r.subject or "[dim](sin asunto)[/dim]",
            llego_de,
            estado,
        )
    console.print(table)
    resumen = (
        f"\n[bold]{len(rows_sorted)}[/bold] registros   "
        f"([yellow]{pend} pendientes[/yellow], [green]{resp} respondidas[/green])"
    )
    if dias_pendientes:
        prom = sum(dias_pendientes) / len(dias_pendientes)
        mayor = max(dias_pendientes)
        resumen += (
            f"\n[bold]Espera promedio:[/bold] {prom:.1f} dias   "
            f"[bold]Mas antigua:[/bold] {_color_dias(mayor)}"
        )
    console.print(resumen)


@app.command()
def show(quote_id: int) -> None:
    """Muestra el detalle completo de una cotizacion."""
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    with Session(engine) as session:
        r = session.get(QuoteRequestRow, quote_id)
        if not r:
            console.print(f"[red]No existe la cotizacion {quote_id}[/red]")
            raise typer.Exit(code=1)

        console.rule(f"Cotizacion #{r.id}")
        console.print(f"[bold]Asunto:[/bold] {r.subject}")
        console.print(f"[bold]Cliente externo:[/bold] {r.external_client or '[dim]no identificado[/dim]'}")
        marca = " [cyan](reenvio interno)[/cyan]" if r.is_forwarded_internal else ""
        console.print(f"[bold]Llego de:[/bold] {r.sender_name} <{r.sender_email}>{marca}")
        console.print(f"[bold]Buzon receptor:[/bold] {r.mailbox}")
        console.print(f"[bold]Recibido:[/bold] {r.received_at}")
        if r.original_source and r.original_source != "mismo_correo":
            console.print(
                f"[bold]Solicitud original:[/bold] {r.original_received_at or '?'} "
                f"[dim]({r.original_source})[/dim]"
            )
            if r.original_snippet:
                console.print(f"  [dim]{r.original_snippet}[/dim]")
        # Fase B: las respuestas viven en quote_related_messages
        respuestas_rows = session.execute(
            select(QuoteRelatedMessage).where(
                QuoteRelatedMessage.quote_request_id == r.id
            ).order_by(QuoteRelatedMessage.received_at)
        ).scalars().all()
        if respuestas_rows:
            console.print(
                f"[bold]Respuestas/RVs vinculadas:[/bold] {len(respuestas_rows)} correo(s)"
            )
            for h in respuestas_rows:
                marca = "[cyan](RV interno)[/cyan] " if h.is_forwarded_internal else ""
                console.print(
                    f"  [dim]- {marca}{(h.subject or '')[:55]} "
                    f"({h.received_at[:10]}, de {h.sender_email})[/dim]"
                )
        if r.chain_source == "backfill_via_RE":
            console.print(
                "  [dim](esta cotizacion se rescato por backfill al detectar una RE)[/dim]"
            )

        if r.group_key:
            hermanos = session.execute(
                select(QuoteRequestRow).where(QuoteRequestRow.group_key == r.group_key)
            ).scalars().all()
            otros = [h for h in hermanos if h.id != r.id]
            dominio, _, ref = r.group_key.partition("|")
            console.print(
                f"[bold]Solicitud consolidada:[/bold] {ref} "
                f"[dim]({len(hermanos)} correos del cliente {dominio})[/dim]"
            )
            for h in otros:
                console.print(f"  [dim]- #{h.id} {h.subject[:50]}[/dim]")
        if r.was_replied:
            console.print(f"[bold green]RESPONDIDA AL CLIENTE[/bold green] el {r.replied_at}")
        else:
            console.print("[bold yellow]PENDIENTE de responder al cliente[/bold yellow]")

        console.print("\n[bold]Resumen de la solicitud:[/bold]")
        console.print(f"  {r.snippet or '[dim](sin texto extraido)[/dim]'}")

        if r.cantidades:
            console.print("\n[bold]Cantidades mencionadas:[/bold]")
            for c in r.cantidades:
                console.print(f"  - {c}")

        if r.adjuntos:
            console.print("\n[bold]Adjuntos:[/bold]")
            for a in r.adjuntos:
                console.print(f"  - {a}")


@app.command()
def cleanup(
    dry_run: bool = typer.Option(
        False, "--dry-run", help="Solo mostrar que se borraria, sin tocar la BD."
    ),
    yes: bool = typer.Option(
        False, "--yes", "-y", help="Saltar confirmacion (no interactivo)."
    ),
    reason: str = typer.Option(
        None, "--reason",
        help="Filtrar por motivo de descarte (ej. cotizacion_saliente, reunion_teams).",
    ),
) -> None:
    """Borra filas de quote_requests cuyo graph_id quedo registrado como
    'descartado' en processed_messages. Util tras afinar reglas para purgar
    falsos positivos legacy (Teams, cadenas salientes, etc.)."""
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    with Session(engine) as session:
        stmt = select(ProcessedMessage.graph_id, ProcessedMessage.reason).where(
            ProcessedMessage.result == "descartado"
        )
        if reason:
            stmt = stmt.where(ProcessedMessage.reason == reason)
        descartados_gids = {row[0]: row[1] for row in session.execute(stmt).all()}

        if not descartados_gids:
            console.print(
                "[green]No hay correos descartados registrados que limpiar.[/green]"
            )
            raise typer.Exit(code=0)

        candidatos = session.execute(
            select(QuoteRequestRow).where(
                QuoteRequestRow.graph_id.in_(descartados_gids.keys())
            )
        ).scalars().all()

    if not candidatos:
        console.print(
            "[green]No hay filas en quote_requests que coincidan con descartados.[/green]"
        )
        raise typer.Exit(code=0)

    # Detectar REs huerfanas que se generarian al borrar
    huerfanas: list[int] = []
    with Session(engine) as session:
        gids_a_borrar = {r.graph_id for r in candidatos}
        for r in session.execute(
            select(QuoteRequestRow).where(
                QuoteRequestRow.original_graph_id.in_(gids_a_borrar)
            )
        ).scalars().all():
            if r.id not in {c.id for c in candidatos}:
                huerfanas.append(r.id)

    table = Table(title="Filas a eliminar de quote_requests", show_lines=True)
    table.add_column("ID", justify="right")
    table.add_column("Buzon", overflow="fold", max_width=24)
    table.add_column("Asunto", overflow="fold", max_width=44)
    table.add_column("Motivo")
    for r in candidatos:
        table.add_row(
            str(r.id), r.mailbox, (r.subject or "")[:55],
            descartados_gids.get(r.graph_id, "?"),
        )
    console.print(table)

    if huerfanas:
        console.print(
            f"\n[yellow]Aviso:[/yellow] al borrar quedarian "
            f"[bold]{len(huerfanas)}[/bold] RE(s) huerfanas: "
            f"{', '.join('#'+str(i) for i in huerfanas[:10])}"
        )

    if dry_run:
        console.print(
            f"\n[dim](dry-run: no se borro nada. "
            f"{len(candidatos)} fila(s) serian eliminadas.)[/dim]"
        )
        raise typer.Exit(code=0)

    if not yes:
        if not typer.confirm(
            f"\nBorrar {len(candidatos)} fila(s)?", default=False
        ):
            console.print("[yellow]Cancelado.[/yellow]")
            raise typer.Exit(code=0)

    with Session(engine) as session:
        ids = [c.id for c in candidatos]
        session.execute(delete(QuoteRequestRow).where(QuoteRequestRow.id.in_(ids)))
        session.commit()
    console.print(
        f"[bold green]Listo:[/bold green] {len(candidatos)} fila(s) eliminada(s) "
        f"de quote_requests."
    )


@app.command("list-threads")
def list_threads() -> None:
    """Lista los HILOS: cada cotizacion (original) con sus respuestas/RVs vinculadas."""
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    from collections import defaultdict
    with Session(engine) as session:
        relacionadas = session.execute(
            select(QuoteRelatedMessage).order_by(QuoteRelatedMessage.received_at)
        ).scalars().all()
        if not relacionadas:
            console.print(
                "[yellow]No hay hilos con respuestas vinculadas todavia.[/yellow]"
            )
            raise typer.Exit(code=0)
        hijos_por_padre: dict[int, list[QuoteRelatedMessage]] = defaultdict(list)
        for r in relacionadas:
            hijos_por_padre[r.quote_request_id].append(r)
        padres = session.execute(
            select(QuoteRequestRow).where(
                QuoteRequestRow.id.in_(list(hijos_por_padre.keys()))
            )
        ).scalars().all()
        padres_map = {p.id: p for p in padres}

    table = Table(title="Hilos: cotizacion original + correos relacionados", show_lines=True)
    table.add_column("ID", justify="right")
    table.add_column("Asunto original", overflow="fold", max_width=42)
    table.add_column("Cliente", overflow="fold", max_width=24)
    table.add_column("Resp.", justify="right")
    table.add_column("Origen", no_wrap=True)
    table.add_column("Estado", no_wrap=True)

    backfill_count = 0
    for pid, hs in sorted(hijos_por_padre.items(), key=lambda kv: -len(kv[1])):
        padre = padres_map.get(pid)
        if not padre:
            continue
        if padre.chain_source == "backfill_via_RE":
            backfill_count += 1
        estado = (
            "[green]RESPONDIDA[/green]" if padre.was_replied
            else "[yellow]PENDIENTE[/yellow]"
        )
        table.add_row(
            f"#{padre.id}",
            (padre.subject or "")[:60],
            padre.external_client or padre.sender_email,
            str(len(hs)),
            padre.chain_source or "",
            estado,
        )
    console.print(table)
    console.print(
        f"\n[bold]{len(hijos_por_padre)}[/bold] hilo(s); "
        f"[bold]{backfill_count}[/bold] cotizacion(es) rescatada(s) por backfill via RE."
    )


@app.command("list-groups")
def list_groups() -> None:
    """Lista las solicitudes CONSOLIDADAS: varios correos de un mismo cliente
    agrupados por numero de documento (escenario 2)."""
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    with Session(engine) as session:
        rows = session.execute(
            select(QuoteRequestRow)
            .where(QuoteRequestRow.group_key.is_not(None))
            .order_by(QuoteRequestRow.group_key)
        ).scalars().all()

    if not rows:
        console.print(
            "[yellow]No hay solicitudes agrupadas "
            "(ningun correo con numero de documento detectado).[/yellow]"
        )
        raise typer.Exit(code=0)

    from collections import defaultdict
    grupos: dict[str, list[QuoteRequestRow]] = defaultdict(list)
    for r in rows:
        grupos[r.group_key].append(r)

    table = Table(title="Solicitudes consolidadas (agrupadas por documento)", show_lines=True)
    table.add_column("Documento", no_wrap=True)
    table.add_column("Cliente", overflow="fold", max_width=26)
    table.add_column("Correos", justify="right")
    table.add_column("Pendientes", justify="right")
    table.add_column("IDs / asuntos", overflow="fold", max_width=50)

    consolidadas = 0
    for key, miembros in grupos.items():
        dominio, _, ref = key.partition("|")
        pend = sum(1 for m in miembros if not m.was_replied)
        if len(miembros) > 1:
            consolidadas += 1
        detalle = "\n".join(f"#{m.id} {(m.subject or '')[:42]}" for m in miembros)
        table.add_row(
            ref,
            dominio,
            str(len(miembros)),
            f"[yellow]{pend}[/yellow]" if pend else "[green]0[/green]",
            detalle,
        )
    console.print(table)
    console.print(
        f"\n[bold]{len(grupos)}[/bold] grupo(s) con documento; "
        f"[bold]{consolidadas}[/bold] de ellos con mas de un correo (consolidados)."
    )


@app.command()
def export(
    output: str = typer.Option(
        None, "--output", "-o", help="Ruta del archivo .xlsx. Por defecto: cotizaciones_pendientes_YYYYMMDD.xlsx"
    ),
    incluir_respondidas: bool = typer.Option(
        False, "--all", help="Incluir tambien las respondidas (por defecto solo pendientes)."
    ),
) -> None:
    """Exporta las cotizaciones PENDIENTES a Excel, ordenadas por antiguedad."""
    settings = load_settings()
    _setup_logging(settings.log_level)
    engine = init_db(settings.database_path)

    with Session(engine) as session:
        stmt = select(QuoteRequestRow).order_by(desc(QuoteRequestRow.received_at))
        if not incluir_respondidas:
            stmt = stmt.where(QuoteRequestRow.was_replied.is_(False))
        rows = session.execute(stmt).scalars().all()

    if not rows:
        console.print("[yellow]No hay cotizaciones para exportar.[/yellow]")
        raise typer.Exit(code=0)

    # Ordenar: pendientes mas antiguas primero
    rows = sorted(
        rows,
        key=lambda r: (r.was_replied, -(_dias_desde(r.received_at) or 0)),
    )

    data = []
    for r in rows:
        dias = _dias_desde(r.received_at)
        data.append({
            "ID": r.id,
            "Dias sin responder": dias if (dias is not None and not r.was_replied) else "",
            "Recibido": (r.received_at or "").replace("T", " ").replace("Z", ""),
            "Cliente externo": r.external_client or "",
            "Asunto": r.subject,
            "Llego de (remitente)": r.sender_email,
            "Es reenvio interno": "SI" if r.is_forwarded_internal else "NO",
            "Nombre remitente": r.sender_name,
            "Estado": "RESPONDIDA AL CLIENTE" if r.was_replied else "PENDIENTE",
            "Respondido en": (r.replied_at or "").replace("T", " ").replace("+00:00", "")
            if r.was_replied else "",
            "Buzon": r.mailbox,
            "Items": r.num_items,
            "Resumen": r.snippet,
            "Cantidades mencionadas": ", ".join(r.cantidades) if r.cantidades else "",
            "Adjuntos": ", ".join(r.adjuntos) if r.adjuntos else "",
        })
    df = pd.DataFrame(data)

    if output is None:
        suffix = "" if incluir_respondidas else "_pendientes"
        output = f"cotizaciones{suffix}_{datetime.now():%Y%m%d_%H%M}.xlsx"
    out_path = Path(output).resolve()

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Cotizaciones", index=False)
        # Auto-ajustar ancho de columnas
        ws = writer.sheets["Cotizaciones"]
        for col_idx, col_name in enumerate(df.columns, start=1):
            max_len = max(
                len(str(col_name)),
                df[col_name].astype(str).map(len).max() if len(df) else 0,
            )
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(
                max_len + 2, 60
            )
        # Resaltar header
        from openpyxl.styles import Font, PatternFill
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
        ws.freeze_panes = "A2"

        # Colorear "Dias sin responder" segun antiguedad
        if "Dias sin responder" in df.columns:
            col_idx = df.columns.get_loc("Dias sin responder") + 1
            fill_red = PatternFill("solid", fgColor="F8CBAD")     # rojo claro
            fill_yellow = PatternFill("solid", fgColor="FFE699")  # amarillo claro
            fill_green = PatternFill("solid", fgColor="C6EFCE")   # verde claro
            for row_idx in range(2, len(df) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                v = cell.value
                if isinstance(v, (int, float)):
                    if v > 3:
                        cell.fill = fill_red
                        cell.font = Font(bold=True)
                    elif v > 1:
                        cell.fill = fill_yellow
                    else:
                        cell.fill = fill_green

    total = len(rows)
    pend = sum(1 for r in rows if not r.was_replied)
    console.print(
        f"\n[bold green]Exportado:[/bold green] {out_path}\n"
        f"  Filas: {total}  ([yellow]{pend} pendientes[/yellow], "
        f"[green]{total - pend} respondidas[/green])"
    )


if __name__ == "__main__":
    app()
